import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm, Twips, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ──────────────────────────────────────────────────────────────
# Design Tokens  –  single source of truth for the entire doc
# ──────────────────────────────────────────────────────────────
FONT_PRIMARY = "Calibri"
FONT_HEADING = "Calibri"

# Strict typographic scale (in points)
SIZE_NAME = 18.5          # Candidate name
SIZE_CONTACT = 9.5        # Contact bar
SIZE_SECTION_HEADING = 11.5 # PROFESSIONAL EXPERIENCE, SKILLS, …
SIZE_ENTRY_TITLE = 11     # Job title / Degree
SIZE_ENTRY_SUB = 10.5     # Company, institution
SIZE_BODY = 10.5          # Bullet text, summary paragraphs
SIZE_META = 10            # Dates, locations

# Colour palette
CLR_NAME = RGBColor(0x1A, 0x1A, 0x2E)       # Deep navy for name
CLR_HEADING = RGBColor(0x2C, 0x3E, 0x50)     # Dark slate for headings
CLR_BODY = RGBColor(0x2D, 0x2D, 0x2D)        # Soft black for body
CLR_ACCENT = RGBColor(0x2C, 0x3E, 0x50)      # Accent (skill cats, contact)
CLR_META = RGBColor(0x55, 0x55, 0x55)         # Grey for dates / locations
CLR_DIVIDER = "2C3E50"                        # Hex for XML borders

# Spacing
MARGIN_TOP = 0.45
MARGIN_BOTTOM = 0.45
MARGIN_LEFT = 0.65
MARGIN_RIGHT = 0.65
LINE_SPACING = 1.15


# ──────────────────────────────────────────────────────────────
# Low‑level helpers
# ──────────────────────────────────────────────────────────────
def _set_margins(doc):
    """Apply tight, consistent page margins."""
    for section in doc.sections:
        section.top_margin = Inches(MARGIN_TOP)
        section.bottom_margin = Inches(MARGIN_BOTTOM)
        section.left_margin = Inches(MARGIN_LEFT)
        section.right_margin = Inches(MARGIN_RIGHT)


def _configure_normal_style(doc):
    """Set the document‑wide Normal style so bare paragraphs inherit sane defaults."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = FONT_PRIMARY
    font.size = Pt(SIZE_BODY)
    font.color.rgb = CLR_BODY
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE_SPACING


def _usable_width_inches(doc) -> float:
    """Return the text‑area width in inches (page − margins)."""
    sec = doc.sections[0]
    return Emu(sec.page_width - sec.left_margin - sec.right_margin).inches


def _add_right_tab(para, doc):
    """Add a right‑aligned tab stop at the right margin."""
    para.paragraph_format.tab_stops.add_tab_stop(
        Inches(_usable_width_inches(doc)),
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )


def _bottom_border(para, weight: int = 4):
    """Attach a thin bottom border to *para*."""
    pPr = para._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="{weight}" w:space="1" w:color="{CLR_DIVIDER}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


# ──────────────────────────────────────────────────────────────
# Document building blocks
# ──────────────────────────────────────────────────────────────
def _add_name(doc, name: str):
    """Large, centred candidate name."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(name)
    run.font.name = FONT_HEADING
    run.font.size = Pt(SIZE_NAME)
    run.font.color.rgb = CLR_NAME
    run.bold = True
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.space_before = Pt(0)


def _add_contact(doc, contact: dict):
    """Single‑line centred contact bar separated by pipes."""
    parts = []
    for key in ("email", "phone", "location", "linkedin", "other"):
        val = contact.get(key)
        if val:
            parts.append(val)
    if not parts:
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("  |  ".join(parts))
    run.font.name = FONT_PRIMARY
    run.font.size = Pt(SIZE_CONTACT)
    run.font.color.rgb = CLR_ACCENT
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.space_before = Pt(0)


def _add_divider(doc):
    """Thin horizontal rule below the contact block."""
    para = doc.add_paragraph()
    _bottom_border(para, weight=6)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.space_before = Pt(0)


def _add_section_heading(doc, title: str):
    """ALL‑CAPS section heading with a subtle underline."""
    para = doc.add_paragraph()
    run = para.add_run(title.upper())
    run.font.name = FONT_HEADING
    run.font.size = Pt(SIZE_SECTION_HEADING)
    run.font.color.rgb = CLR_HEADING
    run.bold = True
    # Add letter‑spacing via XML for a modern feel
    rPr = run._element.get_or_add_rPr()
    spacing_el = parse_xml(f'<w:spacing {nsdecls("w")} w:val="30"/>')
    rPr.append(spacing_el)
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(1)
    _bottom_border(para, weight=4)
    return para


def _add_body_paragraph(doc, text: str):
    """Standard body paragraph (used for summaries, etc.)."""
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = FONT_PRIMARY
    run.font.size = Pt(SIZE_BODY)
    run.font.color.rgb = CLR_BODY
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    para.paragraph_format.line_spacing = LINE_SPACING
    return para


def _add_bullet(doc, text: str):
    """Bullet‑pointed body text."""
    para = doc.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    run.font.name = FONT_PRIMARY
    run.font.size = Pt(SIZE_BODY)
    run.font.color.rgb = CLR_BODY
    para.paragraph_format.space_before = Pt(0.5)
    para.paragraph_format.space_after = Pt(0.5)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    para.paragraph_format.line_spacing = LINE_SPACING
    return para


def _add_experience_header(doc, title: str, company: str, location: str, dates: str):
    """Job title | Company, Location ―(tab)→ Dates."""
    para = doc.add_paragraph()
    _add_right_tab(para, doc)

    # Title (bold, slightly larger)
    title_run = para.add_run(title)
    title_run.font.name = FONT_PRIMARY
    title_run.font.size = Pt(SIZE_ENTRY_TITLE)
    title_run.font.color.rgb = CLR_BODY
    title_run.bold = True

    # Company
    if company:
        sep = para.add_run("  |  ")
        sep.font.name = FONT_PRIMARY
        sep.font.size = Pt(SIZE_ENTRY_SUB)
        sep.font.color.rgb = CLR_META
        co_run = para.add_run(company)
        co_run.font.name = FONT_PRIMARY
        co_run.font.size = Pt(SIZE_ENTRY_SUB)
        co_run.font.color.rgb = CLR_BODY

    # Location
    if location:
        loc_run = para.add_run(", " + location)
        loc_run.font.name = FONT_PRIMARY
        loc_run.font.size = Pt(SIZE_META)
        loc_run.font.color.rgb = CLR_META
        loc_run.italic = True

    # Dates (right‑aligned via tab)
    if dates:
        dates_run = para.add_run("\t" + dates)
        dates_run.font.name = FONT_PRIMARY
        dates_run.font.size = Pt(SIZE_META)
        dates_run.font.color.rgb = CLR_META
        dates_run.italic = True

    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    para.paragraph_format.line_spacing = LINE_SPACING


def _add_education_entry(doc, edu: dict):
    """Degree  –  Institution, Location ―(tab)→ Dates."""
    para = doc.add_paragraph()
    _add_right_tab(para, doc)

    deg_run = para.add_run(edu.get("degree", ""))
    deg_run.font.name = FONT_PRIMARY
    deg_run.font.size = Pt(SIZE_ENTRY_TITLE)
    deg_run.font.color.rgb = CLR_BODY
    deg_run.bold = True

    institution = edu.get("institution", "")
    if institution:
        sep = para.add_run("  –  ")
        sep.font.name = FONT_PRIMARY
        sep.font.size = Pt(SIZE_ENTRY_SUB)
        sep.font.color.rgb = CLR_META
        inst_run = para.add_run(institution)
        inst_run.font.name = FONT_PRIMARY
        inst_run.font.size = Pt(SIZE_ENTRY_SUB)
        inst_run.font.color.rgb = CLR_BODY

    location = edu.get("location", "")
    if location:
        loc_run = para.add_run(", " + location)
        loc_run.font.name = FONT_PRIMARY
        loc_run.font.size = Pt(SIZE_META)
        loc_run.font.color.rgb = CLR_META
        loc_run.italic = True

    dates = edu.get("dates", "")
    if dates:
        dates_run = para.add_run("\t" + dates)
        dates_run.font.name = FONT_PRIMARY
        dates_run.font.size = Pt(SIZE_META)
        dates_run.font.color.rgb = CLR_META
        dates_run.italic = True

    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    para.paragraph_format.line_spacing = LINE_SPACING

    details = edu.get("details", "")
    if details:
        _add_bullet(doc, details)


def _add_skills_section(doc, skills: dict):
    """Category: skill, skill, skill — compact two-column feel."""
    _add_section_heading(doc, "Skills")
    for category, skill_list in skills.items():
        para = doc.add_paragraph()
        cat_run = para.add_run(category + ":  ")
        cat_run.font.name = FONT_PRIMARY
        cat_run.font.size = Pt(SIZE_BODY)
        cat_run.font.color.rgb = CLR_ACCENT
        cat_run.bold = True
        skills_run = para.add_run(", ".join(skill_list))
        skills_run.font.name = FONT_PRIMARY
        skills_run.font.size = Pt(SIZE_BODY)
        skills_run.font.color.rgb = CLR_BODY
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        para.paragraph_format.line_spacing = LINE_SPACING


def _add_certifications_section(doc, certifications: list):
    """Certification list with issuer and date."""
    _add_section_heading(doc, "Certifications")
    for cert in certifications:
        text = cert.get("name", "")
        issuer = cert.get("issuer", "")
        date = cert.get("date", "")
        if issuer:
            text += " – " + issuer
        if date:
            text += " (" + date + ")"
        _add_bullet(doc, text)


# ──────────────────────────────────────────────────────────────
# Public API
# ──────────────────────────────────────────────────────────────
def generate_docx(tailored_cv: dict, resume_styles: dict, output_path: str) -> str:
    """Build a professionally‑formatted DOCX from the structured CV dict.

    *resume_styles* is accepted for API compatibility with app.py but is
    intentionally **not used** for formatting.  The previous approach of
    re‑applying extracted source styles produced inconsistent font sizes
    and spacing.  This version uses a self‑contained design system instead.
    """
    doc = Document()

    # ── Page setup ──────────────────────────────────────────
    _set_margins(doc)
    _configure_normal_style(doc)

    # ── Header block ────────────────────────────────────────
    _add_name(doc, tailored_cv.get("name", "Candidate"))
    _add_contact(doc, tailored_cv.get("contact", {}))
    _add_divider(doc)

    # ── Professional Summary ────────────────────────────────
    summary = tailored_cv.get("professional_summary", "")
    if summary:
        _add_section_heading(doc, "Professional Summary")
        _add_body_paragraph(doc, summary)

    # ── Skills ──────────────────────────────────────────────
    skills = tailored_cv.get("skills", {})
    if skills:
        _add_skills_section(doc, skills)

    # ── Professional Experience ─────────────────────────────
    experience = tailored_cv.get("experience", [])
    if experience:
        _add_section_heading(doc, "Professional Experience")
        for exp in experience:
            _add_experience_header(
                doc,
                exp.get("title", ""),
                exp.get("company", ""),
                exp.get("location", ""),
                exp.get("dates", ""),
            )
            for highlight in exp.get("highlights", []):
                _add_bullet(doc, highlight)

    # ── Education ───────────────────────────────────────────
    education = tailored_cv.get("education", [])
    if education:
        _add_section_heading(doc, "Education")
        for edu in education:
            _add_education_entry(doc, edu)

    # ── Certifications ──────────────────────────────────────
    certifications = tailored_cv.get("certifications", [])
    if certifications:
        _add_certifications_section(doc, certifications)

    # ── Additional / custom sections ────────────────────────
    additional = tailored_cv.get("additional_sections", {})
    for section_name, content in additional.items():
        _add_section_heading(doc, section_name)
        if isinstance(content, list):
            for item in content:
                _add_bullet(doc, str(item))
        elif isinstance(content, str):
            _add_body_paragraph(doc, content)

    # ── Save ────────────────────────────────────────────────
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def convert_to_pdf(docx_path: str, pdf_path: str = None) -> str:
    import pythoncom
    from docx2pdf import convert

    if pdf_path is None:
        pdf_path = docx_path.replace(".docx", ".pdf")

    # Streamlit runs handlers in worker threads where COM is not
    # initialised.  We must bracket the Word automation call.
    pythoncom.CoInitialize()
    try:
        convert(docx_path, pdf_path)
    finally:
        pythoncom.CoUninitialize()
    return pdf_path
