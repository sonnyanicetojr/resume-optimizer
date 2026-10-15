import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH


def extract_text(docx_path: str) -> str:
    doc = Document(docx_path)
    full_text = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            full_text.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text)


def _extract_font_info(run) -> dict:
    font_info = {}
    if run.font.name:
        font_info["name"] = run.font.name
    if run.font.size:
        font_info["size"] = run.font.size.pt
    if run.font.bold is not None:
        font_info["bold"] = run.font.bold
    if run.font.italic is not None:
        font_info["italic"] = run.font.italic
    if run.font.color and run.font.color.rgb:
        font_info["color"] = str(run.font.color.rgb)
    if run.font.underline is not None:
        font_info["underline"] = run.font.underline
    return font_info


def _extract_paragraph_style(para) -> dict:
    style = {
        "alignment": str(para.alignment) if para.alignment is not None else None,
        "space_before": para.paragraph_format.space_before.pt if para.paragraph_format.space_before else None,
        "space_after": para.paragraph_format.space_after.pt if para.paragraph_format.space_after else None,
        "line_spacing": para.paragraph_format.line_spacing,
        "left_indent": para.paragraph_format.left_indent.inches if para.paragraph_format.left_indent else None,
    }
    if para.runs:
        style["runs"] = [_extract_font_info(run) for run in para.runs]
    return style


def extract_styles(docx_path: str) -> dict:
    doc = Document(docx_path)

    section_props = {}
    for section in doc.sections:
        section_props = {
            "page_width": section.page_width.inches,
            "page_height": section.page_height.inches,
            "left_margin": section.left_margin.inches,
            "right_margin": section.right_margin.inches,
            "top_margin": section.top_margin.inches,
            "bottom_margin": section.bottom_margin.inches,
        }
        break

    style_map = {}
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else "Normal"
        if style_name not in style_map:
            style_map[style_name] = _extract_paragraph_style(para)

    return {
        "sections": section_props,
        "paragraph_styles": style_map,
    }


def parse_resume(docx_path: str) -> dict:
    text = extract_text(docx_path)
    styles = extract_styles(docx_path)
    return {
        "text": text,
        "styles": styles,
    }


def extract_name_from_resume(docx_path: str) -> str:
    doc = Document(docx_path)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and len(text.split()) >= 2:
            likely_name = True
            for word in text.split():
                if word.lower() in ["email", "phone", "http", "www", "@", "cell", "tel"]:
                    likely_name = False
                    break
            if likely_name:
                return text
    return "Candidate"